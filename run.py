"""
run.py — INTERN / ENTRY-LEVEL Resume & Cover Letter Maker: ONE command from input Excel to finished PDFs + updated Excel.

    python run.py                                   # processes every new .xlsx in Input\\
    python run.py Input\\Intern_Jobs_2026-09-26.xlsx  # one file
    python run.py <file> --model claude-haiku-4-5-20251001 --workers 6
    python run.py <file> --redo                     # regenerate an already-processed file

Per file:
  1. FILE      Input\\<M_D_YYYY>\\original\\<file>.xlsx   (date from filename, else today)
  2. ADDRESS   complete_address_section  (exact city -> nearest same-state -> Houston)
  3. ATS       ats  (URL table + deep fingerprint of company sites)
  4. JD        jd   (Jobright structured JSON, fetched once, cached in the Excel)
  5. GENERATE  Output\\<M_D_YYYY>\\<Company>_<Title>\\Bhanu_Prakash_Bathini.pdf (config.RESUME_PAGES page)
                                                   + Bhanu_Prakash_Bathini_Cover_Letter.pdf (config.COVER_PAGES page)
               LLM (2 calls/job) -> JSON -> render.py: validate, block fabrication + seniority, LaTeX, auto-fit, verify
  6. WRITE     updated_resume_section + coverletter  ->  Input\\<M_D_YYYY>\\updated\\<file>_updated.xlsx
  7. REPORT    logs\\<M_D_YYYY>_<file>.json
"""
import argparse, concurrent.futures as cf, glob, json, os, re, shutil, sys, time
from copy import copy
from datetime import date

APP = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, APP); sys.path.insert(0, os.path.join(APP, "engine"))
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
import config, address, jd_fetch
from engine import ats_detect, prompts, render, llm

# ------------------------------------------------------------------ helpers
def slug(s): return re.sub(r"[^A-Za-z0-9]+", "_", str(s)).strip("_")[:45]

def day_from(filename):
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", filename)
    if m: return f"{int(m.group(2))}_{int(m.group(3))}_{m.group(1)}"
    t = date.today(); return f"{t.month}_{t.day}_{t.year}"

def link(cell):
    """Excel cells here are hyperlinks with display text 'Open posting' / 'Jobright' — take the target."""
    if cell.hyperlink and cell.hyperlink.target: return cell.hyperlink.target
    v = str(cell.value or "").strip()
    return v if v.startswith("http") else ""

def col(ws, *names, exact=False):
    hdr = [str(c.value).strip() if c.value else "" for c in ws[1]]
    for n in names:
        if exact:
            if n in hdr: return hdr.index(n) + 1
        else:
            low = [h.lower() for h in hdr]
            if n.lower() in low: return low.index(n.lower()) + 1
    return None

def ensure_columns(ws):
    """Append config.NEW_COLUMNS after the last header (skip ones that already exist). -> {name: col}"""
    out = {}
    hf = copy(ws.cell(row=1, column=1).font)
    for name in config.NEW_COLUMNS:
        c = col(ws, name, exact=True)          # exact: the input may already carry 'ATS' from the classifier
        if not c:
            c = ws.max_column + 1
            ws.cell(row=1, column=c, value=name).font = hf
            ws.column_dimensions[get_column_letter(c)].width = 60 if name == "jd" else 45
        out[name] = c
    return out

def is_entry_level(title, level, job_type=""):
    t = f"{title} {level} {job_type}".lower()
    return any(w in t for w in config.ENTRY_LEVEL_WORDS)

def log(msg): print(time.strftime("%H:%M:%S"), msg, flush=True)

def save_wb(wb, path):
    """Save; if Excel has the file open, save next to it as *_PENDING.xlsx and keep going."""
    try:
        wb.save(path); return path
    except PermissionError:
        alt = path.replace(".xlsx", "_PENDING.xlsx"); wb.save(alt)
        log(f"    !! {os.path.basename(path)} is OPEN in Excel -> saved {os.path.basename(alt)}. Close Excel and rerun to merge.")
        return alt

# ------------------------------------------------------------------ steps
def stage(xlsx_path, redo=False):
    """Steps 1-3 + folder pre-staging. Returns (day, updated_xlsx, jobs)."""
    fname = os.path.basename(xlsx_path); day = day_from(fname)
    orig_dir = os.path.join(config.INPUT_DIR, day, "original"); upd_dir = os.path.join(config.INPUT_DIR, day, "updated")
    os.makedirs(orig_dir, exist_ok=True); os.makedirs(upd_dir, exist_ok=True)
    dst = os.path.join(orig_dir, fname)
    if os.path.abspath(xlsx_path) != os.path.abspath(dst): shutil.move(xlsx_path, dst)
    upd_xlsx = os.path.join(upd_dir, os.path.splitext(fname)[0] + "_updated.xlsx")
    resuming = os.path.exists(upd_xlsx) and not redo
    log(f"[1] {'resuming' if resuming else 'filed'} -> {upd_xlsx if resuming else dst}")

    # resume from the updated workbook so cached JDs / addresses survive; fresh runs start from the original
    pending = upd_xlsx.replace(".xlsx", "_PENDING.xlsx")
    src_wb = pending if (resuming and os.path.exists(pending)) else (upd_xlsx if resuming else dst)
    wb = openpyxl.load_workbook(src_wb); ws = wb[wb.sheetnames[0]]
    if src_wb == pending:
        try: os.remove(pending)
        except OSError: pass
    c_title, c_co = col(ws, "Job Title") or 2, col(ws, "Company") or 3
    c_loc, c_lvl  = col(ws, "Location") or 6, col(ws, "Level") or 9
    c_type        = col(ws, "Job Type") or 8
    c_orig, c_jr  = col(ws, "Original Job Posting Link") or 5, col(ws, "Jobright Link") or 14
    cols = ensure_columns(ws)
    out_root = os.path.join(config.OUTPUT_DIR, day); os.makedirs(out_root, exist_ok=True)

    jobs, stats, seen, not_entry = [], {}, {}, []
    for r in range(2, ws.max_row + 1):
        title = ws.cell(r, c_title).value
        if not title: continue
        company = str(ws.cell(r, c_co).value or "Unknown").strip()
        level, jtype = str(ws.cell(r, c_lvl).value or ""), str(ws.cell(r, c_type).value or "")
        if not is_entry_level(title, level, jtype): not_entry.append((ws.cell(r, 1).value, company, level))
        addr, how = address.resolve(ws.cell(r, c_loc).value); stats[how] = stats.get(how, 0) + 1
        ws.cell(r, cols["complete_address_section"], value=addr)
        o_url, j_url = link(ws.cell(r, c_orig)), link(ws.cell(r, c_jr))
        ws.cell(r, cols["ats"], value=ats_detect.detect(o_url, company, deep=True))
        folder = f"{slug(company)}_{slug(title)}"; key = folder.lower()          # Windows: case-insensitive
        seen[key] = seen.get(key, 0) + 1
        if seen[key] > 1: folder += f"_{seen[key]}"
        jobs.append({"row": r, "num": ws.cell(r, 1).value, "company": company, "title": str(title).strip(),
                     "level": level, "job_type": jtype, "city_st": address.city_st(addr),
                     "address": addr, "original": o_url, "jobright": j_url, "folder": folder,
                     "dir": os.path.join(out_root, folder), "jd": str(ws.cell(r, cols["jd"]).value or "")})
    save_wb(wb, upd_xlsx)
    log(f"[2] addresses {stats}  [3] ats written  -> {upd_xlsx}")
    if not_entry:
        log(f"    note: {len(not_entry)} row(s) do not look intern/entry-level (still processed, entry-level tone applied): {not_entry[:6]}")
    for j in jobs:
        os.makedirs(j["dir"], exist_ok=True)
        json.dump({k: j[k] for k in ("row", "num", "company", "title", "level", "job_type", "city_st", "folder", "jd")},
                  open(os.path.join(j["dir"], "job.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    return day, upd_xlsx, jobs

def fetch_jds(jobs, upd_xlsx, workers=2):
    """Step 4: fetch every missing JD (2 threads + 1 s spacing — Jobright rate-limits bursts),
    write the `jd` column + job.json."""
    # manual override: drop a jd_manual.txt into the job folder when the link is dead (403 / removed)
    for j in jobs:
        mp = os.path.join(j["dir"], "jd_manual.txt")
        if os.path.exists(mp) and (not j["jd"] or j["jd"].startswith("FETCH_FAILED")):
            j["jd"] = open(mp, encoding="utf-8").read().strip(); j["jd_source"] = "manual"; j["_write"] = True
    todo = [j for j in jobs if not j["jd"] or j["jd"].startswith("FETCH_FAILED")]
    if todo:
        def _one(j):
            time.sleep(1.0)
            return jd_fetch.fetch(j["jobright"], j["original"], j["company"])
        with cf.ThreadPoolExecutor(workers) as ex:
            for j, (txt, src) in zip(todo, ex.map(_one, todo)):
                j["jd"], j["jd_source"] = txt, src
    todo = todo + [j for j in jobs if j.get("_write") and j not in todo]
    if todo:
        wb = openpyxl.load_workbook(upd_xlsx); ws = wb[wb.sheetnames[0]]; c_jd = col(ws, "jd", exact=True)
        for j in todo:
            c = ws.cell(j["row"], c_jd, value=j["jd"]); c.alignment = Alignment(wrap_text=True, vertical="top")
            meta = json.load(open(os.path.join(j["dir"], "job.json"), encoding="utf-8")); meta["jd"] = j["jd"]
            json.dump(meta, open(os.path.join(j["dir"], "job.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
        save_wb(wb, upd_xlsx)
    ok = sum(1 for j in jobs if j["jd"] and not j["jd"].startswith("FETCH_FAILED"))
    log(f"[4] jd {ok}/{len(jobs)} cached" + (f"  FAILED: {[j['num'] for j in jobs if j['jd'].startswith('FETCH_FAILED')]}" if ok < len(jobs) else ""))
    return ok

def generate_one(j, args, date_str):
    """Step 5 for one job: 2 LLM calls -> render.build_job (validate/compile/fit/verify)."""
    t0 = time.time()
    if not j["jd"] or j["jd"].startswith("FETCH_FAILED"):
        return {"folder": j["folder"], "ok": False, "errors": ["no JD"]}
    jd, co, ti, cs = j["jd"], j["company"], j["title"], j["city_st"]
    lv = " / ".join(x for x in (j["level"], j.get("job_type", "")) if x)
    spec, errs = llm.json_call(args, prompts.build_messages(jd, co, ti, cs, lv), prompts.SYSTEM_RESUME,
                               lambda o: render.validate(render.normalize(o), cs))
    if spec is None: return {"folder": j["folder"], "ok": False, "stage": "resume-llm", "errors": errs}
    def cl_check(o):
        if not isinstance(o, dict): return ["cover letter must be a JSON object with addr,p1,p2,p3,p4"]
        miss = [k for k in ("p1", "p2", "p3", "p4") if not str(o.get(k, "")).strip()]
        if miss: return ["missing paragraphs: " + ", ".join(miss)]
        low = json.dumps(o).lower()
        fab = [b for b in render.BLOCKLIST if re.search(rf"(?<![a-z0-9]){re.escape(b)}(?![a-z0-9])", low)]
        errs = ["remove these technologies (not in BASE_RESUME): " + ", ".join(fab)] if fab else []
        sen = [w for w in render.SENIOR_WORDS if re.search(rf"\b{re.escape(w)}\b", low)]
        sen += [f"years claim '{m.group(0)}'" for m in render.YEARS_RE.finditer(" ".join(str(o.get(k, "")) for k in ("p1", "p2", "p3", "p4")))
                if m.group(0).lower() not in render._BASE_TEXT]
        if sen: errs.append("entry-level tone: remove " + ", ".join(sen))
        return errs
    cl, errs = llm.json_call(args, prompts.build_cl_messages(jd, co, ti, cs, lv), prompts.SYSTEM_CL, cl_check)
    if cl is None: return {"folder": j["folder"], "ok": False, "stage": "cl-llm", "errors": errs}
    res = render.build_job(j["dir"], spec, cl, date_str)
    res["seconds"] = round(time.time() - t0, 1)
    if not res["ok"]: res["errors"] = (res["resume"].get("errors") or []) + (res["cover_letter"].get("errors") or [])
    return res

def verify(j):
    rp, cp = os.path.join(j["dir"], config.RESUME_PDF), os.path.join(j["dir"], config.COVER_PDF)
    try:
        return (os.path.exists(rp) and os.path.getsize(rp) > 15360 and len(PdfReader(rp).pages) == config.RESUME_PAGES
                and os.path.exists(cp) and len(PdfReader(cp).pages) == config.COVER_PAGES)
    except Exception:
        return False

def write_paths(jobs, upd_xlsx):
    wb = openpyxl.load_workbook(upd_xlsx); ws = wb[wb.sheetnames[0]]
    c_res, c_cl = col(ws, "updated_resume_section", exact=True), col(ws, "coverletter", exact=True)
    n = 0
    for j in jobs:
        if verify(j):
            ws.cell(j["row"], c_res, value=os.path.join(j["dir"], config.RESUME_PDF))
            ws.cell(j["row"], c_cl,  value=os.path.join(j["dir"], config.COVER_PDF)); n += 1
    save_wb(wb, upd_xlsx); return n

def cleanup(jobs):
    for j in jobs:
        for f in os.listdir(j["dir"]):
            if f.endswith((".aux", ".log", ".out", ".synctex.gz")) or f in ("main.pdf", "coverletter.pdf"):
                try: os.remove(os.path.join(j["dir"], f))
                except OSError: pass

# ------------------------------------------------------------------ main
def process_file(xlsx_path, args):
    t0 = time.time()
    day, upd_xlsx, jobs = stage(xlsx_path, args.redo)
    if not args.redo and all(verify(j) for j in jobs):
        log(f"ALREADY PROCESSED: {os.path.basename(xlsx_path)} -> {upd_xlsx}  (use --redo to regenerate)"); return
    fetch_jds(jobs, upd_xlsx)
    t = date.today(); date_str = t.strftime("%B ") + str(t.day) + t.strftime(", %Y")
    todo = jobs if args.redo else [j for j in jobs if not verify(j)]
    log(f"[5] generating {len(todo)} jobs | model={args.model} | workers={args.workers} | resume={config.RESUME_PAGES}p cover={config.COVER_PAGES}p")
    results = []
    with cf.ThreadPoolExecutor(args.workers) as ex:
        futs = {ex.submit(generate_one, j, args, date_str): j for j in todo}
        for i, f in enumerate(cf.as_completed(futs), 1):
            r = f.result(); results.append(r)
            tag = "OK" if r["ok"] else "FAIL " + "; ".join(r.get("errors", []))[:100]
            log(f"    [{i:>3}/{len(todo)}] {r['folder'][:48]:<48} {tag}  ({r.get('seconds','?')}s)")
    # one retry pass for failures (new sampling usually fixes a bad spec)
    failed = [j for j in todo if not verify(j)]
    if failed:
        log(f"    retrying {len(failed)} failed job(s)")
        with cf.ThreadPoolExecutor(min(args.workers, len(failed))) as ex:
            for r in ex.map(lambda j: generate_one(j, args, date_str), failed):
                log(f"    retry {r['folder'][:48]:<48} {'OK' if r['ok'] else 'FAIL ' + '; '.join(r.get('errors', []))[:100]}")
    n = write_paths(jobs, upd_xlsx); cleanup(jobs)
    good = [j for j in jobs if verify(j)]; bad = [j for j in jobs if not verify(j)]
    report = {"file": os.path.basename(xlsx_path), "day": day, "jobs": len(jobs), "generated": n,
              "failed": [{"num": j["num"], "company": j["company"], "title": j["title"]} for j in bad],
              "model": args.model, "resume_pages": config.RESUME_PAGES, "minutes": round((time.time() - t0) / 60, 1),
              "excel": upd_xlsx, "output_dir": os.path.join(config.OUTPUT_DIR, day)}
    os.makedirs(config.LOG_DIR, exist_ok=True)
    json.dump(report, open(os.path.join(config.LOG_DIR, f"{day}_{os.path.splitext(os.path.basename(xlsx_path))[0]}.json"), "w"), indent=1)
    log(f"[6] DONE {n}/{len(jobs)} jobs in {report['minutes']} min -> {upd_xlsx}")
    if bad: log(f"    FAILED: {[(j['num'], j['company']) for j in bad]}")
    return report

def main():
    ap = argparse.ArgumentParser(description="Resume & Cover Letter Maker — intern / entry-level")
    ap.add_argument("xlsx", nargs="?", help="input Excel (default: every new .xlsx directly in Input\\)")
    ap.add_argument("--model", default=config.LLM_MODEL)
    ap.add_argument("--base-url", default=config.LLM_BASE_URL)
    ap.add_argument("--api-key", default=config.LLM_API_KEY)
    ap.add_argument("--workers", type=int, default=config.LLM_WORKERS)
    ap.add_argument("--max-tokens", type=int, default=config.LLM_MAX_TOKENS)
    ap.add_argument("--redo", action="store_true")
    args = ap.parse_args()
    files = [args.xlsx] if args.xlsx else [p for p in glob.glob(os.path.join(config.INPUT_DIR, "*.xlsx")) if not os.path.basename(p).startswith("~$")]
    if not files: log("nothing to do: no .xlsx in Input\\"); return
    for f in files:
        try: process_file(f, args)
        except PermissionError as e: log(f"PermissionError ({e}): close the Excel file in the Excel app and rerun")

if __name__ == "__main__":
    main()
