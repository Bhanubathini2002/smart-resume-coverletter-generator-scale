"""
ats_detect.py — which ATS / job board / company site a job posting lives on.

Tier 1 (instant, offline): the ORIGINAL job posting URL's host+path against ~60 known
    ATS vendors, ~12 job boards, and company-name heuristics.
Tier 2 (optional, --deep): for links that resolve to a company website, GET the page
    (12 s, follows redirects) and fingerprint the HTML for embedded ATS vendors
    (Phenom, Workday, Greenhouse embed, iCIMS, SmartRecruiters, ...) → "Workday (via company site)".

Column value (header `ats`) is one of:
    <Vendor>                      Greenhouse, Lever, Workday, Ashby, iCIMS, ...
    <Vendor> (via company site)   company careers page that embeds/redirects to a vendor
    Company website               employer-hosted apply, no known vendor detected
    <Board> (job board)           LinkedIn, Dice, ZipRecruiter, Indeed, ... — apply is elsewhere
    Jobright only                 no original link in the row
    Unknown (<host>)              third-party host not in the table

CLI:
    python engine/ats_detect.py <M_D_YYYY> [--deep] [--workers 12]        # writes `ats` into Input/<day>/updated/*_updated.xlsx
    python engine/ats_detect.py --url https://boards.greenhouse.io/x/jobs/1  # one-off
    python engine/ats_detect.py --xlsx path.xlsx [--deep]                   # any file with an 'Original Job Posting Link' column
"""
import argparse, glob, os, re, sys, ssl, urllib.request, urllib.error
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urlparse

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---- Tier 1: host / path patterns ------------------------------------------
# (regex on "host/path", first match wins — order matters only for overlaps)
ATS_PATTERNS = [
    (r"greenhouse\.io|grnh\.se",                       "Greenhouse"),
    (r"lever\.co",                                     "Lever"),
    (r"myworkdayjobs\.com|myworkdaysite\.com|workday\.com", "Workday"),
    (r"ashbyhq\.com",                                  "Ashby"),
    (r"icims\.com",                                    "iCIMS"),
    (r"smartrecruiters\.com",                          "SmartRecruiters"),
    (r"taleo\.net",                                    "Taleo"),
    (r"oraclecloud\.com",                              "Oracle Recruiting Cloud"),
    (r"successfactors\.com|jobs\.sap\.com",            "SAP SuccessFactors"),
    (r"brassring\.com|kenexa",                         "IBM Kenexa BrassRing"),
    (r"eightfold\.ai",                                 "Eightfold"),
    (r"avature\.net",                                  "Avature"),
    (r"csod\.com",                                     "Cornerstone OnDemand"),
    (r"rippling\.com",                                 "Rippling"),
    (r"paylocity\.com",                                "Paylocity"),
    (r"paycomonline\.net",                             "Paycom"),
    (r"ukg\.net|ultipro\.com",                         "UKG"),
    (r"dayforcehcm\.com",                              "Dayforce"),
    (r"workforcenow\.adp\.com|adp\.com",               "ADP"),
    (r"workable\.com",                                 "Workable"),
    (r"breezy\.hr",                                    "Breezy HR"),
    (r"applytojob\.com|jazz\.co|jazzhr\.com",          "JazzHR"),
    (r"bamboohr\.com",                                 "BambooHR"),
    (r"jobvite\.com",                                  "Jobvite"),
    (r"jobs\.gem\.com|gem\.com",                       "Gem"),
    (r"recruiterflow\.com",                            "Recruiterflow"),
    (r"loxo\.co",                                      "Loxo"),
    (r"catsone\.com",                                  "CATS"),
    (r"hrmdirect\.com",                                "HRM Direct"),
    (r"trinethire\.com",                               "TriNet Hire"),
    (r"pcrecruiter\.net",                              "PCRecruiter"),
    (r"njoyn\.com",                                    "Njoyn"),
    (r"contacthr\.com",                                "ContactHR"),
    (r"careers-page\.com|manatal\.com",                "Manatal"),
    (r"teamtailor\.com",                               "Teamtailor"),
    (r"recruitee\.com",                                "Recruitee"),
    (r"personio\.(com|de)",                            "Personio"),
    (r"pinpointhq\.com",                               "Pinpoint"),
    (r"dover\.com",                                    "Dover"),
    (r"clearcompany\.com",                             "ClearCompany"),
    (r"applicantpro\.com",                             "ApplicantPro"),
    (r"governmentjobs\.com",                           "NEOGOV"),
    (r"usajobs\.gov",                                  "USAJOBS"),
    (r"comeet\.com",                                   "Comeet"),
    (r"freshteam\.com",                                "Freshteam"),
    (r"zohorecruit\.com",                              "Zoho Recruit"),
    (r"homerun\.co",                                   "Homerun"),
    (r"polymer\.co",                                   "Polymer"),
    (r"phenompeople\.com|phenom\.com",                 "Phenom"),
    (r"radancy\.com|tmp\.com",                         "Radancy"),
    (r"jibeapply\.com",                                "iCIMS (Jibe)"),
    (r"hirebridge\.com",                               "Hirebridge"),
    (r"silkroad\.com",                                 "SilkRoad"),
    (r"crelate\.com",                                  "Crelate"),
    (r"bullhornreach\.com|bullhorn\.com",              "Bullhorn"),
    (r"jobdiva\.com",                                  "JobDiva"),
    (r"ceipal\.com",                                   "Ceipal"),
    (r"hr\.cloud\.sap|\.jobs\.sap\b",                  "SAP SuccessFactors"),
    (r"recruitingbypaycor\.com|paycor\.com",           "Paycor"),
    (r"careers/jobdetail\?|/en_us/careers/jobdetail",  "Radancy"),          # IBM / Deloitte / Slalom pattern
    (r"/psp/[a-z0-9]*erecruit|hrs_hram|/psc/",         "PeopleSoft"),
    (r"\.jobs/job/[a-z0-9-]+/\d+",                     "SAP SuccessFactors"), # komatsu.jobs/job/.../35901-en_US
    (r"/job/[a-z0-9%-]+/\d{8,}/?$",                    "SAP SuccessFactors"), # kiewitcareers…/1410311300/
    (r"ycombinator\.com/companies/.*/jobs",            "Y Combinator Work at a Startup"),
    (r"ibegin\.tcsapps\.com|ibegin\.tcs\.com",           "TCS iBegin"),
]
BOARD_PATTERNS = [
    (r"linkedin\.com",                  "LinkedIn"),
    (r"dice\.com",                      "Dice"),
    (r"ziprecruiter\.com",              "ZipRecruiter"),
    (r"indeed\.com",                    "Indeed"),
    (r"glassdoor\.com",                 "Glassdoor"),
    (r"builtin\.com|builtin[a-z]*\.com","Built In"),
    (r"simplyhired\.com",               "SimplyHired"),
    (r"wellfound\.com|angel\.co",       "Wellfound"),
    (r"welcometothejungle\.com|otta\.com", "Welcome to the Jungle"),
    (r"monster\.com",                   "Monster"),
    (r"careerbuilder\.com",             "CareerBuilder"),
    (r"consulting\.us",                 "consulting.us"),
    (r"jobright\.ai",                   "Jobright"),
    (r"google\.com/search",             "Google search"),
]
# HTML fingerprints for Tier 2 (regex on lowercase page source); first match wins
HTML_FINGERPRINTS = [
    (r"greenhouse\.io|grnhse|gh_jid",           "Greenhouse"),
    (r"myworkdayjobs|myworkdaysite|workdayjobs|wd\d\.myworkday", "Workday"),
    (r"lever\.co/",                              "Lever"),
    (r"ashbyhq\.com",                            "Ashby"),
    (r"icims\.com|jibeapply",                    "iCIMS"),
    (r"smartrecruiters\.com",                    "SmartRecruiters"),
    (r"phenompeople|phenom\.com|window\.phapp|\"phenom\"|ph-page|pcs-", "Phenom"),
    (r"taleo\.net",                              "Taleo"),
    (r"oraclecloud\.com",                        "Oracle Recruiting Cloud"),
    (r"successfactors",                          "SAP SuccessFactors"),
    (r"brassring|kenexa",                        "IBM Kenexa BrassRing"),
    (r"eightfold\.ai",                           "Eightfold"),
    (r"avature",                                 "Avature"),
    (r"csod\.com",                               "Cornerstone OnDemand"),
    (r"jobvite",                                 "Jobvite"),
    (r"bamboohr",                                "BambooHR"),
    (r"workable\.com",                           "Workable"),
    (r"recruitee",                               "Recruitee"),
    (r"teamtailor",                              "Teamtailor"),
    (r"personio",                                "Personio"),
    (r"rippling\.com",                           "Rippling"),
    (r"paylocity",                               "Paylocity"),
    (r"ultipro|ukg\.net",                        "UKG"),
    (r"dayforcehcm",                             "Dayforce"),
    (r"radancy|tmpw\.",                          "Radancy"),
    (r"symphonytalent",                          "Symphony Talent"),
    (r"beamery",                                 "Beamery"),
]

def _host_path(url):
    p = urlparse(url.strip())
    return (p.netloc or "").lower().lstrip("www."), (p.path or "").lower()

def _company_tokens(company):
    toks = [t for t in re.split(r"[^a-z0-9]+", str(company or "").lower()) if len(t) >= 3]
    stop = {"inc", "llc", "ltd", "corp", "corporation", "company", "the", "and", "group", "technologies", "technology", "systems", "solutions", "labs", "usa"}
    return [t for t in toks if t not in stop]

def detect_url(url, company=""):
    """Tier 1. Returns (label, kind) — kind in {'ats','board','company','none','unknown'}."""
    if not url or not str(url).startswith("http"):
        return "Jobright only", "none"
    host, path = _host_path(str(url))
    hp = host + path
    for pat, name in BOARD_PATTERNS:
        if re.search(pat, hp):
            if name == "Jobright": return "Jobright only", "none"
            return f"{name} (job board)", "board"
    for pat, name in ATS_PATTERNS:
        if name and re.search(pat, hp):
            return name, "ats"
    # company-name match or careers-style host/path → employer-hosted
    flat = host.replace("-", "").replace(".", "")
    if any(t in flat for t in _company_tokens(company)):
        return "Company website", "company"
    if re.search(r"^(careers?|jobs?|apply|talent|recruit|hiring|work)[.-]|[.-](careers?|jobs)\.|/careers?/|/jobs?/|/job/|\.jobs$", host + "/" + path.lstrip("/")):
        return "Company website", "company"
    return f"Unknown ({host})", "unknown"

# ---- Tier 2: fetch + fingerprint ----------------------------------------------
_CTX = ssl.create_default_context(); _CTX.check_hostname = False; _CTX.verify_mode = ssl.CERT_NONE
_UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"

def fingerprint(url, timeout=12):
    """Follow redirects, then match final host (Tier 1 table) and HTML fingerprints.
    Returns vendor name or None."""
    try:
        req = urllib.request.Request(url, headers={"User-Agent": _UA, "Accept": "text/html,*/*;q=0.8", "Accept-Language": "en-US,en;q=0.9"})
        with urllib.request.urlopen(req, timeout=timeout, context=_CTX) as r:
            final = r.geturl(); html = r.read(600_000).decode("utf-8", "ignore").lower()
    except Exception:
        return None
    fh, fp = _host_path(final)
    for pat, name in ATS_PATTERNS:
        if name and re.search(pat, fh + fp): return name
    for pat, name in HTML_FINGERPRINTS:
        if re.search(pat, html): return name
    return None

def detect(url, company="", deep=False):
    label, kind = detect_url(url, company)
    if deep and kind in ("company", "unknown"):
        v = fingerprint(url)
        if v: return f"{v} (via company site)" if kind == "company" else v
    return label

# ---- Excel integration ----------------------------------------------------------
def _find_col(ws, *names):
    low = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    for n in names:
        if n.lower() in low: return low.index(n.lower()) + 1
    return None

def write_excel(xlsx, deep=False, workers=12):
    import openpyxl
    from openpyxl.utils import get_column_letter
    from copy import copy
    wb = openpyxl.load_workbook(xlsx); ws = wb[wb.sheetnames[0]]
    c_url = _find_col(ws, "Original Job Posting Link", "original_link", "apply_url", "url")
    c_co  = _find_col(ws, "Company") or 3
    if not c_url: sys.exit("no 'Original Job Posting Link' column in " + xlsx)
    c_ats = _find_col(ws, "ats")
    if not c_ats:
        c_ats = ws.max_column + 1
        ws.cell(row=1, column=c_ats, value="ats").font = copy(ws.cell(row=1, column=1).font)
        ws.column_dimensions[get_column_letter(c_ats)].width = 30
    rows = [(r, str(ws.cell(row=r, column=c_url).value or ""), ws.cell(row=r, column=c_co).value)
            for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=2).value]
    with ThreadPoolExecutor(max_workers=workers if deep else 1) as ex:
        labels = list(ex.map(lambda t: detect(t[1], t[2], deep), rows))
    stats = {}
    for (r, _, _), lab in zip(rows, labels):
        ws.cell(row=r, column=c_ats, value=lab)
        key = lab.split(" (")[0] if "(" in lab and not lab.startswith("Unknown") else lab
        stats[key] = stats.get(key, 0) + 1
    try: wb.save(xlsx)
    except PermissionError: sys.exit("PermissionError: close the Excel file in the Excel app and rerun")
    return stats, len(rows)

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("day", nargs="?", help="M_D_YYYY -> Input/<day>/updated/*_updated.xlsx")
    ap.add_argument("--xlsx"); ap.add_argument("--url"); ap.add_argument("--company", default="")
    ap.add_argument("--deep", action="store_true", help="fetch company-site links and fingerprint the embedded ATS")
    ap.add_argument("--workers", type=int, default=12)
    a = ap.parse_args()
    if a.url:
        print(detect(a.url, a.company, a.deep)); sys.exit()
    xlsx = a.xlsx
    if not xlsx and a.day:
        c = [p for p in glob.glob(os.path.join(ROOT, "Input", a.day, "updated", "*_updated.xlsx")) if not os.path.basename(p).startswith("~$")]
        if not c: sys.exit("no *_updated.xlsx for " + a.day)
        xlsx = c[0]
    if not xlsx: ap.error("give <day>, --xlsx or --url")
    stats, n = write_excel(xlsx, a.deep, a.workers)
    print(f"ats: {n} rows -> {xlsx}")
    for k, v in sorted(stats.items(), key=lambda kv: -kv[1]): print(f"  {v:4d}  {k}")
